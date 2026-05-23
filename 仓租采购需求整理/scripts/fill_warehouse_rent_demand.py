#!/usr/bin/env python3
import argparse
import copy
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator


DEFAULT_SHEET = "26年618临租"
DATA_COLUMNS = range(1, 18)


def normalize_name(value):
    if value is None:
        return ""
    text = str(value).strip()
    text = re.sub(r"\s+", "", text)
    text = text.replace("仓库", "")
    text = re.sub(r"[（(][^）)]*轮[）)]", "", text)
    return text


def parse_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value
    text = str(value).strip().replace("/", "-").replace(".", "-")
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    raise ValueError(f"无法识别日期：{value!r}")


def as_number(value, field):
    if value in (None, ""):
        return None
    try:
        return float(value)
    except Exception as exc:
        raise ValueError(f"{field} 必须是数字，当前为 {value!r}") from exc


def is_data_row(ws, row):
    return bool(ws.cell(row, 3).value and ws.cell(row, 4).value and ws.cell(row, 5).value)


def find_last_contiguous_data_row(ws):
    last = 1
    for row in range(2, ws.max_row + 1):
        if is_data_row(ws, row):
            last = row
            continue
        if last > 1 and row - last > 3:
            break
    return last


def build_f_lookup(ws, last_data_row):
    lookup = {}
    conflicts = {}
    current_name = None
    for row in range(2, last_data_row + 1):
        name = ws.cell(row, 2).value
        if name:
            current_name = name
        key = normalize_name(name or current_name)
        price = ws.cell(row, 6).value
        if not key or price in (None, "") or isinstance(price, str) and price.startswith("="):
            continue
        old = lookup.get(key)
        if old is not None and float(old) != float(price):
            conflicts.setdefault(key, sorted({float(old), float(price)}))
        lookup[key] = float(price)
    return lookup, conflicts


def copy_row_format(ws, source_row, target_row):
    for col in DATA_COLUMNS:
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy.copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy.copy(src.alignment)
        if src.border:
            dst.border = copy.copy(src.border)
        if src.fill:
            dst.fill = copy.copy(src.fill)
        if src.protection:
            dst.protection = copy.copy(src.protection)
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def translated_formula(ws, source_cell, target_cell, fallback):
    value = ws[source_cell].value
    if isinstance(value, str) and value.startswith("="):
        return Translator(value, origin=source_cell).translate_formula(target_cell)
    return fallback


def validate_record(record, row_num, f_lookup, conflicts):
    missing = []
    warehouse = record.get("warehouse")
    if not warehouse:
        missing.append("仓库")
    if as_number(record.get("area"), "area") is None:
        missing.append("临租面积")
    if parse_date(record.get("start_date")) is None:
        missing.append("起租日期")
    if parse_date(record.get("end_date")) is None:
        missing.append("到期日期")

    key = normalize_name(warehouse)
    main_price = as_number(record.get("main_contract_price"), "main_contract_price")
    if main_price is None:
        if key in conflicts:
            missing.append(f"主合同单价F（历史匹配值冲突：{conflicts[key]}）")
        elif key not in f_lookup:
            missing.append("主合同单价F")

    g = as_number(record.get("temp_unit_price"), "temp_unit_price")
    m = as_number(record.get("total_rent"), "total_rent")
    if g is None and m is None:
        missing.append("临租单价G或总租金M")

    i = as_number(record.get("rent_days_with_free"), "rent_days_with_free")
    j = as_number(record.get("free_days"), "free_days")
    if i is None and j is None:
        missing.append("免租期J或含免租租赁天数I")

    if missing:
        label = warehouse or f"第{row_num}条"
        raise ValueError(f"{label} 缺少：{', '.join(missing)}")


def fill_row(ws, row, record, f_lookup, source_row):
    warehouse = record.get("warehouse")
    area = as_number(record.get("area"), "area")
    start = parse_date(record.get("start_date"))
    end = parse_date(record.get("end_date"))
    key = normalize_name(warehouse)
    main_price = as_number(record.get("main_contract_price"), "main_contract_price")
    if main_price is None:
        main_price = f_lookup[key]

    temp_price = as_number(record.get("temp_unit_price"), "temp_unit_price")
    total_rent = as_number(record.get("total_rent"), "total_rent")
    free_days = as_number(record.get("free_days"), "free_days")
    rent_days_with_free = as_number(record.get("rent_days_with_free"), "rent_days_with_free")

    ws.cell(row, 1).value = record.get("sequence")
    ws.cell(row, 2).value = warehouse
    ws.cell(row, 3).value = area
    ws.cell(row, 4).value = start
    ws.cell(row, 5).value = end
    ws.cell(row, 6).value = main_price

    ws.cell(row, 8).value = record.get("rent_days_without_free") or f"=E{row}-D{row}+1"
    if free_days is None:
        free_days = 0
    ws.cell(row, 10).value = free_days
    if rent_days_with_free is None:
        ws.cell(row, 9).value = f"=H{row}+J{row}"
    else:
        ws.cell(row, 9).value = rent_days_with_free

    h_value = record.get("rent_days_without_free")
    if h_value is None:
        h_value = (end - start).days + 1

    if rent_days_with_free is None:
        i_value = h_value + free_days
    else:
        i_value = rent_days_with_free

    if temp_price is None:
        calculated_g = total_rent / i_value / area
        ws.cell(row, 13).value = total_rent
        ws.cell(row, 7).value = f"=M{row}/I{row}/C{row}"
        gm_action = "G由M/I/C计算"
    elif total_rent is None:
        calculated_g = temp_price
        total_rent = area * temp_price * i_value
        ws.cell(row, 7).value = temp_price
        ws.cell(row, 13).value = f"=C{row}*G{row}*I{row}"
        gm_action = "M由C*G*I计算"
    else:
        calculated_g = temp_price
        ws.cell(row, 7).value = temp_price
        ws.cell(row, 13).value = total_rent
        gm_action = "G和M均采用用户输入"

    actual_price = calculated_g * h_value / i_value
    discount = (main_price - actual_price) / main_price

    ws.cell(row, 11).value = translated_formula(ws, f"K{source_row}", f"K{row}", f"=G{row}*H{row}/I{row}")
    ws.cell(row, 12).value = translated_formula(ws, f"L{source_row}", f"L{row}", f"=(F{row}-K{row})/F{row}")
    ws.cell(row, 14).value = translated_formula(ws, f"N{source_row}", f"N{row}", f"=F{row}*I{row}*C{row}")
    ws.cell(row, 15).value = translated_formula(ws, f"O{source_row}", f"O{row}", f"=N{row}-M{row}")
    ws.cell(row, 16).value = translated_formula(ws, f"P{source_row}", f"P{row}", f"=O{row}/N{row}")
    ws.cell(row, 17).value = translated_formula(ws, f"Q{source_row}", f"Q{row}", f"=O{row}/M{row}")

    return {
        "row": row,
        "warehouse": warehouse,
        "main_contract_price": main_price,
        "gm_action": gm_action,
        "email_row": {
            "仓库": warehouse,
            "临租面积": area,
            "起租日期": start.strftime("%Y-%m-%d"),
            "到期日期": end.strftime("%Y-%m-%d"),
            "主合同单价": main_price,
            "临租单价": calculated_g,
            "租赁时间_不含免租期": h_value,
            "租赁时间_含免租期": i_value,
            "免租期": free_days,
            "实际单价": actual_price,
            "降幅": discount,
            "总租金": total_rent,
        },
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--records", required=True)
    parser.add_argument("--output")
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    args = parser.parse_args()

    workbook_path = Path(args.workbook)
    records_path = Path(args.records)
    output_path = Path(args.output) if args.output else workbook_path

    records = json.loads(records_path.read_text(encoding="utf-8"))
    if not isinstance(records, list) or not records:
        raise ValueError("records.json 必须是非空数组")

    wb = load_workbook(workbook_path)
    ws = wb[args.sheet] if args.sheet in wb.sheetnames else wb.active
    last_data_row = find_last_contiguous_data_row(ws)
    f_lookup, conflicts = build_f_lookup(ws, last_data_row)

    for idx, record in enumerate(records, 1):
        validate_record(record, idx, f_lookup, conflicts)

    insert_at = last_data_row + 1
    ws.insert_rows(insert_at, amount=len(records))

    results = []
    source_row = last_data_row
    for offset, record in enumerate(records):
        row = insert_at + offset
        copy_row_format(ws, source_row, row)
        results.append(fill_row(ws, row, record, f_lookup, source_row))

    if args.output:
        output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(json.dumps({"output": str(output_path), "inserted": results}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
